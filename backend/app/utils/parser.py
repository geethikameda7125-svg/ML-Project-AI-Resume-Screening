import re
import pymupdf as fitz  # PyMuPDF
import docx
from typing import Tuple, Dict, Any

class DocumentParseError(Exception):
    """Custom exception raised when document parsing fails or validation checks fail."""
    pass

def sanitize_pii(text: str) -> str:
    """
    Masks unnecessary personal identifiable information (emails, phone numbers, web addresses)
    to protect candidate privacy while preserving technical skills and context.
    """
    if not text:
        return ""
    
    # Mask emails
    text = re.sub(r'[a-zA-Z0-9._%+-]+@[a-zA-Z0-9.-]+\.[a-zA-Z]{2,}', '[CONFIDENTIAL_EMAIL]', text)
    
    # Mask phone numbers (standard international and local formats)
    text = re.sub(r'(\+?\d{1,3}[-.\s]?)?(\(?\d{3}\)?[-.\s]?)?\d{3}[-.\s]?\d{4}', '[CONFIDENTIAL_PHONE]', text)
    
    # Mask URLs/LinkedIn/Portfolios if needed, preserving general text structure
    text = re.sub(r'https?://[^\s]+', '[CONFIDENTIAL_URL]', text)
    
    return text

def extract_text_from_pdf(file_bytes: bytes) -> str:
    """Extracts text from PDF bytes using PyMuPDF (fitz)."""
    try:
        doc = fitz.open(stream=file_bytes, filetype="pdf")
        if doc.page_count == 0:
            raise DocumentParseError("The uploaded PDF document contains 0 pages.")
        
        extracted_pages = []
        for page_num in range(doc.page_count):
            page = doc.load_page(page_num)
            text = page.get_text("text")
            if text and text.strip():
                extracted_pages.append(text.strip())
        
        full_text = "\n\n".join(extracted_pages)
        if not full_text.strip():
            raise DocumentParseError(
                "Unable to extract text from PDF. The document may be scanned, image-only, or encrypted. "
                "Please upload a searchable text-based PDF or DOCX resume."
            )
        return full_text
    except Exception as e:
        if isinstance(e, DocumentParseError):
            raise e
        raise DocumentParseError(f"Error parsing PDF file: {str(e)}")

def extract_text_from_docx(file_bytes: bytes) -> str:
    """Extracts text from DOCX bytes using python-docx."""
    try:
        import io
        doc = docx.Document(io.BytesIO(file_bytes))
        paragraphs = [p.text.strip() for p in doc.paragraphs if p.text and p.text.strip()]
        
        # Also extract table text if any
        for table in doc.tables:
            for row in table.rows:
                row_text = " | ".join(cell.text.strip() for cell in row.cells if cell.text.strip())
                if row_text:
                    paragraphs.append(row_text)
                    
        full_text = "\n".join(paragraphs)
        if not full_text.strip():
            raise DocumentParseError("The uploaded DOCX document is empty or unreadable.")
        return full_text
    except Exception as e:
        if isinstance(e, DocumentParseError):
            raise e
        raise DocumentParseError(f"Error parsing DOCX file: {str(e)}")

def extract_text_from_txt(file_bytes: bytes) -> str:
    """Extracts text from UTF-8 / ASCII plain text bytes."""
    try:
        text = file_bytes.decode("utf-8", errors="ignore").strip()
        if not text:
            raise DocumentParseError("The uploaded text file is empty.")
        return text
    except Exception as e:
        raise DocumentParseError(f"Error reading TXT file: {str(e)}")

def parse_resume_document(filename: str, file_bytes: bytes) -> Dict[str, Any]:
    """
    Master parser routing files based on extension, extracting text, sanitizing PII,
    and returning metadata.
    """
    if not file_bytes:
        raise DocumentParseError("Uploaded file is empty (0 bytes).")
    
    ext = filename.lower().split('.')[-1]
    
    if ext == "pdf":
        raw_text = extract_text_from_pdf(file_bytes)
    elif ext in ["docx", "doc"]:
        raw_text = extract_text_from_docx(file_bytes)
    elif ext == "txt":
        raw_text = extract_text_from_txt(file_bytes)
    else:
        raise DocumentParseError(f"Unsupported file format '.{ext}'. Please upload a PDF, DOCX, or TXT document.")
    
    # Check minimum meaningful length (e.g. at least 30 characters)
    clean_raw = raw_text.strip()
    if len(clean_raw) < 30:
        raise DocumentParseError("Document contains insufficient text content for meaningful resume analysis.")
        
    sanitized = sanitize_pii(clean_raw)
    
    return {
        "filename": filename,
        "file_type": ext,
        "raw_character_count": len(clean_raw),
        "sanitized_text": sanitized
    }
