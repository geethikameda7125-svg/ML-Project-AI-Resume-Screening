import sys
import os
import pptx
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml import parse_xml

def create_humanized_ppt():
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    blank_slide_layout = prs.slide_layouts[6]

    # Theme Colors (Formal White Background + Deep Royal Blue Headings)
    COLOR_BG = RGBColor(255, 255, 255)          # White Background (#FFFFFF)
    COLOR_CARD = RGBColor(248, 250, 252)        # Light Slate Card (#F8FAFC)
    COLOR_CARD_BORDER = RGBColor(203, 213, 225) # Slate Border (#CBD5E1)
    
    COLOR_BLUE_HEADING = RGBColor(30, 58, 138)  # Deep Royal Blue (#1E3A8A)
    COLOR_BLUE_SUB = RGBColor(37, 99, 235)      # Bright Accent Blue (#2563EB)
    COLOR_DARK_TEXT = RGBColor(30, 41, 59)       # Charcoal Text (#1E293B)
    COLOR_SUBTEXT = RGBColor(71, 85, 105)       # Slate Gray (#475569)
    COLOR_TEAL = RGBColor(13, 148, 136)        # Accent Teal (#0D9488)

    def set_slide_background(slide):
        background = slide.background
        fill = background.fill
        fill.solid()
        fill.fore_color.rgb = COLOR_BG

    def add_header(slide, title_text, category_text="SHORT-TERM INDUSTRY INTERNSHIP REVIEW"):
        tag_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.4), Inches(11.733), Inches(0.35))
        tf_tag = tag_box.text_frame
        p_tag = tf_tag.paragraphs[0]
        p_tag.text = category_text.upper()
        p_tag.font.size = Pt(10)
        p_tag.font.bold = True
        p_tag.font.color.rgb = COLOR_BLUE_SUB

        title_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.7), Inches(11.733), Inches(0.75))
        tf_title = title_box.text_frame
        p_title = tf_title.paragraphs[0]
        p_title.text = title_text
        p_title.font.size = Pt(22)
        p_title.font.bold = True
        p_title.font.color.rgb = COLOR_BLUE_HEADING

    def apply_thick_black_borders(table):
        for row in table.rows:
            for cell in row.cells:
                tcPr = cell._tc.get_or_add_tcPr()
                for side in ['lnL', 'lnR', 'lnT', 'lnB']:
                    existing = tcPr.find(f'{{http://schemas.openxmlformats.org/drawingml/2006/main}}{side}')
                    if existing is not None:
                        tcPr.remove(existing)
                    ln_xml = f'<a:{side} xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main" w="31750" cmpd="s"><a:solidFill><a:srgbClr val="000000"/></a:solidFill></a:{side}>'
                    tcPr.append(parse_xml(ln_xml))

    def add_notes(slide, notes_text):
        notes_slide = slide.notes_slide
        tf = notes_slide.notes_text_frame
        tf.text = notes_text

    # ==========================================
    # SLIDE 1: TITLE SLIDE (BIG BOLD TITLE AT TOP)
    # ==========================================
    s1 = prs.slides.add_slide(blank_slide_layout)
    set_slide_background(s1)

    tag_box1 = s1.shapes.add_textbox(Inches(0.8), Inches(0.4), Inches(11.733), Inches(0.4))
    tf_t1 = tag_box1.text_frame
    p = tf_t1.paragraphs[0]
    p.text = "SHORT-TERM INDUSTRY INTERNSHIP REVIEW PRESENTATION"
    p.font.size = Pt(11)
    p.font.bold = True
    p.font.color.rgb = COLOR_BLUE_SUB

    title_box1 = s1.shapes.add_textbox(Inches(0.8), Inches(0.8), Inches(11.733), Inches(1.4))
    tf_title1 = title_box1.text_frame
    tf_title1.word_wrap = True
    p_title1 = tf_title1.paragraphs[0]
    p_title1.text = "AI-Based Resume Screening and Job Matching System"
    p_title1.font.size = Pt(32)
    p_title1.font.bold = True
    p_title1.font.color.rgb = COLOR_BLUE_HEADING

    sub_box1 = s1.shapes.add_textbox(Inches(0.8), Inches(2.1), Inches(11.733), Inches(0.5))
    tf_sub1 = sub_box1.text_frame
    p_sub1 = tf_sub1.paragraphs[0]
    p_sub1.text = "Real-Time Machine Learning & Natural Language Processing Project"
    p_sub1.font.size = Pt(16)
    p_sub1.font.bold = True
    p_sub1.font.color.rgb = COLOR_TEAL

    c_main = s1.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(2.7), Inches(11.733), Inches(4.3))
    c_main.fill.solid()
    c_main.fill.fore_color.rgb = COLOR_CARD
    c_main.line.color.rgb = COLOR_BLUE_SUB
    c_main.line.width = Pt(1.5)

    tf_m = c_main.text_frame
    tf_m.word_wrap = True
    tf_m.margin_left = Inches(0.4)
    tf_m.margin_top = Inches(0.3)

    p = tf_m.paragraphs[0]
    p.text = "INSTITUTION & DEPARTMENT"
    p.font.size = Pt(11)
    p.font.bold = True
    p.font.color.rgb = COLOR_BLUE_SUB

    p = tf_m.add_paragraph()
    p.text = "G. Pulla Reddy Engineering College (Autonomous), Kurnool"
    p.font.size = Pt(14)
    p.font.bold = True
    p.font.color.rgb = COLOR_DARK_TEXT

    p = tf_m.add_paragraph()
    p.text = "Department of Computer Science and Engineering"
    p.font.size = Pt(12)
    p.font.color.rgb = COLOR_SUBTEXT
    p.space_after = Pt(16)

    p = tf_m.add_paragraph()
    p.text = "PRESENTER & REGISTRATION DETAILS"
    p.font.size = Pt(11)
    p.font.bold = True
    p.font.color.rgb = COLOR_BLUE_SUB

    p = tf_m.add_paragraph()
    p.text = "Student Name: Meda Geethika  |  Register Number: 239X1A05D6"
    p.font.size = Pt(13)
    p.font.bold = True
    p.font.color.rgb = COLOR_DARK_TEXT

    p = tf_m.add_paragraph()
    p.text = "Degree: B.Tech – Computer Science and Engineering"
    p.font.size = Pt(11.5)
    p.font.color.rgb = COLOR_SUBTEXT
    p.space_after = Pt(16)

    p = tf_m.add_paragraph()
    p.text = "INTERNSHIP ORGANIZATION & FACULTY GUIDE"
    p.font.size = Pt(11)
    p.font.bold = True
    p.font.color.rgb = COLOR_BLUE_SUB

    p = tf_m.add_paragraph()
    p.text = "Organization: InternPe  |  Domain: Artificial Intelligence / Machine Learning  |  Duration: 8 Weeks (04-05-2026 to 28-06-2026)"
    p.font.size = Pt(11.5)
    p.font.color.rgb = COLOR_DARK_TEXT

    p = tf_m.add_paragraph()
    p.text = "Faculty Guide: Smt. Y. Supriya Reddy, Assistant Professor, Dept. of CSE"
    p.font.size = Pt(12)
    p.font.bold = True
    p.font.color.rgb = COLOR_BLUE_HEADING

    add_notes(s1, "Respected evaluators and my faculty guide Smt. Y. Supriya Reddy ma'am, good morning. I am Meda Geethika (239X1A05D6), B.Tech CSE student at G. Pulla Reddy Engineering College. Today I present my short-term internship review on 'AI-Based Resume Screening and Job Matching System' completed at InternPe in the AI/ML domain.")

    # ==========================================
    # SLIDE 2: REGISTRATION DETAILS (TABLE 1 OF 4)
    # ==========================================
    s2 = prs.slides.add_slide(blank_slide_layout)
    set_slide_background(s2)
    add_header(s2, "Student & Internship Registration Details")

    table_shape = s2.shapes.add_table(10, 2, Inches(0.8), Inches(1.5), Inches(11.733), Inches(5.3))
    table = table_shape.table
    table.columns[0].width = Inches(3.5)
    table.columns[1].width = Inches(8.233)

    details = [
        ("Student Name", "Meda Geethika"),
        ("Register Number", "239X1A05D6"),
        ("Degree & Branch", "B.Tech – Computer Science and Engineering"),
        ("Institution", "G. Pulla Reddy Engineering College (Autonomous), Kurnool"),
        ("Internship Organization", "InternPe"),
        ("Internship Domain", "Artificial Intelligence / Machine Learning"),
        ("Internship Title", "Real Time Machine Learning Project Using Python"),
        ("Duration & Dates", "8 Weeks (04-05-2026 to 28-06-2026)"),
        ("Mode of Internship", "Virtual / Remote"),
        ("Faculty Guide", "Smt. Y. Supriya Reddy, Assistant Professor, Dept. of CSE")
    ]

    for idx, (label, val) in enumerate(details):
        cell_lbl = table.cell(idx, 0)
        cell_val = table.cell(idx, 1)

        cell_lbl.fill.solid()
        cell_lbl.fill.fore_color.rgb = COLOR_CARD
        cell_val.fill.solid()
        cell_val.fill.fore_color.rgb = COLOR_BG

        p0 = cell_lbl.text_frame.paragraphs[0]
        p0.text = label
        p0.font.size = Pt(11)
        p0.font.bold = True
        p0.font.color.rgb = COLOR_BLUE_HEADING

        p1 = cell_val.text_frame.paragraphs[0]
        p1.text = val
        p1.font.size = Pt(11)
        p1.font.color.rgb = COLOR_DARK_TEXT

    apply_thick_black_borders(table)
    add_notes(s2, "This table details my official internship registration parameters. Completed over 8 weeks with InternPe under academic guidance of Smt. Y. Supriya Reddy ma'am.")

    # ==========================================
    # SLIDE 3: ABOUT INTERNPE & ROLE
    # ==========================================
    s3 = prs.slides.add_slide(blank_slide_layout)
    set_slide_background(s3)
    add_header(s3, "About Organization & Internship Role")

    c_l = s3.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.5), Inches(5.6), Inches(5.3))
    c_l.fill.solid()
    c_l.fill.fore_color.rgb = COLOR_CARD
    c_l.line.color.rgb = COLOR_CARD_BORDER
    tf_l = c_l.text_frame
    tf_l.word_wrap = True
    tf_l.margin_left = Inches(0.3)
    tf_l.margin_top = Inches(0.3)

    p = tf_l.paragraphs[0]
    p.text = "About InternPe Organization"
    p.font.size = Pt(16)
    p.font.bold = True
    p.font.color.rgb = COLOR_BLUE_HEADING
    p.space_after = Pt(14)

    bullets_l = [
        "Technology platform focusing on practical, project-based engineering skill enhancement.",
        "Offers remote internships in Software Development, Data Science, AI, and Machine Learning.",
        "Emphasizes task-oriented learning, Python framework application, and full-stack integration."
    ]
    for b in bullets_l:
        p = tf_l.add_paragraph()
        p.text = "• " + b
        p.font.size = Pt(11.5)
        p.font.color.rgb = COLOR_DARK_TEXT
        p.space_after = Pt(10)

    c_r = s3.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(6.8), Inches(1.5), Inches(5.733), Inches(5.3))
    c_r.fill.solid()
    c_r.fill.fore_color.rgb = COLOR_CARD
    c_r.line.color.rgb = COLOR_CARD_BORDER
    tf_r = c_r.text_frame
    tf_r.word_wrap = True
    tf_r.margin_left = Inches(0.3)
    tf_r.margin_top = Inches(0.3)

    p = tf_r.paragraphs[0]
    p.text = "My Role & Internship Responsibilities"
    p.font.size = Pt(16)
    p.font.bold = True
    p.font.color.rgb = COLOR_BLUE_HEADING
    p.space_after = Pt(14)

    bullets_r = [
        "AI/ML Engineer Intern: Designed end-to-end NLP data pipelines and machine learning algorithms.",
        "Built document extraction modules using PyMuPDF and python-docx for PDF/DOCX resume parsing.",
        "Implemented TF-IDF feature engineering, Cosine Similarity matching engine, and FastAPI REST endpoints.",
        "Developed React.js interactive frontend dashboard and SQLite history persistence."
    ]
    for b in bullets_r:
        p = tf_r.add_paragraph()
        p.text = "• " + b
        p.font.size = Pt(11.5)
        p.font.color.rgb = COLOR_DARK_TEXT
        p.space_after = Pt(10)

    add_notes(s3, "InternPe provided hands-on project exposure. My role encompassed the complete ML pipeline from text parsing to web interface deployment.")

    # ==========================================
    # SLIDE 4: OBJECTIVES & MILESTONES
    # ==========================================
    s4 = prs.slides.add_slide(blank_slide_layout)
    set_slide_background(s4)
    add_header(s4, "Internship Learning Objectives & Technical Goals")

    objs = [
        ("1. Master ML Pipeline", "Understand complete end-to-end workflow from unstructured text parsing to TF-IDF feature vector creation."),
        ("2. NLP Text Processing", "Apply text cleaning, tokenization, stop-word filtering, and synonym dictionary mapping."),
        ("3. Mathematical Matching", "Implement Cosine Similarity to quantitatively measure vector alignment between resume and job description."),
        ("4. Gap & Roadmap Engine", "Extract missing candidate skills and map actionable learning recommendations for career growth."),
        ("5. REST API Integration", "Build FastAPI backend endpoints for document parsing, scoring, and history logging."),
        ("6. Production UI & DB", "Create responsive React.js frontend with SQLite database storage for analysis history.")
    ]

    for idx, (title, desc) in enumerate(objs):
        row = idx // 3
        col = idx % 3
        left = Inches(0.8 + col * 3.95)
        top = Inches(1.5 + row * 2.7)

        card = s4.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, Inches(3.8), Inches(2.4))
        card.fill.solid()
        card.fill.fore_color.rgb = COLOR_CARD
        card.line.color.rgb = COLOR_CARD_BORDER

        tf = card.text_frame
        tf.word_wrap = True
        tf.margin_left = Inches(0.25)
        tf.margin_top = Inches(0.2)

        p = tf.paragraphs[0]
        p.text = title
        p.font.size = Pt(13)
        p.font.bold = True
        p.font.color.rgb = COLOR_BLUE_HEADING

        p2 = tf.add_paragraph()
        p2.text = desc
        p2.font.size = Pt(10.5)
        p2.font.color.rgb = COLOR_DARK_TEXT
        p2.space_before = Pt(6)

    add_notes(s4, "These 6 objectives structured my 8-week engineering goal, ensuring theoretical concepts were translated into a working application.")

    # ==========================================
    # SLIDE 5: PROBLEM STATEMENT & FLOWCHART
    # ==========================================
    s5 = prs.slides.add_slide(blank_slide_layout)
    set_slide_background(s5)
    add_header(s5, "Problem Statement & Process Comparison")

    c_trad = s5.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.5), Inches(5.6), Inches(5.3))
    c_trad.fill.solid()
    c_trad.fill.fore_color.rgb = COLOR_CARD
    c_trad.line.color.rgb = COLOR_CARD_BORDER
    tf_t = c_trad.text_frame
    tf_t.word_wrap = True
    tf_t.margin_left = Inches(0.3)
    tf_t.margin_top = Inches(0.3)

    p = tf_t.paragraphs[0]
    p.text = "Traditional Manual Screening"
    p.font.size = Pt(15)
    p.font.bold = True
    p.font.color.rgb = COLOR_BLUE_HEADING

    bullets_t = [
        "High Time Consumption: Manual reading of hundreds of resumes takes days.",
        "Keyword Bias & Oversight: Critical technical skills easily overlooked in long text.",
        "No Candidate Feedback: Applicants receive no insights into missing skills or gap areas."
    ]
    for b in bullets_t:
        p = tf_t.add_paragraph()
        p.text = "• " + b
        p.font.size = Pt(11)
        p.font.color.rgb = COLOR_DARK_TEXT
        p.space_after = Pt(10)

    c_prop = s5.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(6.8), Inches(1.5), Inches(5.733), Inches(5.3))
    c_prop.fill.solid()
    c_prop.fill.fore_color.rgb = COLOR_CARD
    c_prop.line.color.rgb = COLOR_BLUE_SUB
    c_prop.line.width = Pt(1.5)
    tf_p = c_prop.text_frame
    tf_p.word_wrap = True
    tf_p.margin_left = Inches(0.3)
    tf_p.margin_top = Inches(0.3)

    p = tf_p.paragraphs[0]
    p.text = "Proposed NLP/ML Solution"
    p.font.size = Pt(15)
    p.font.bold = True
    p.font.color.rgb = COLOR_BLUE_HEADING

    bullets_p = [
        "Automated Multi-Format Parsing: Instantly extracts text from PDF, DOCX, TXT.",
        "TF-IDF & Cosine Vector Match: Mathematical similarity calculation unaffected by formatting.",
        "Transparent Skill Gap Insights: Identifies matched vs missing skills and guides candidate learning."
    ]
    for b in bullets_p:
        p = tf_p.add_paragraph()
        p.text = "• " + b
        p.font.size = Pt(11)
        p.font.color.rgb = COLOR_DARK_TEXT
        p.space_after = Pt(10)

    add_notes(s5, "This slide compares traditional manual screening against our proposed ML/NLP approach, emphasizing objective matching and skill gap guidance.")

    # ==========================================
    # SLIDE 6: PROPOSED SYSTEM FLOWCHART
    # ==========================================
    s6 = prs.slides.add_slide(blank_slide_layout)
    set_slide_background(s6)
    add_header(s6, "Proposed System Flowchart & High-Level Workflow")

    steps_s6 = [
        ("1. Input Document", "PDF / DOCX / TXT Resume Upload"),
        ("2. Text Extraction", "PyMuPDF & python-docx Parsing"),
        ("3. Preprocessing", "Cleaning, Stop-words, Normalization"),
        ("4. Feature Vectorization", "TF-IDF Vector Representation"),
        ("5. Cosine Similarity", "Vector Alignment Calculation"),
        ("6. Skill Matching", "Matched & Missing Skill Extraction"),
        ("7. Action Roadmap", "Targeted Learning Guidance")
    ]

    for idx, (step_title, step_desc) in enumerate(steps_s6):
        left = Inches(0.8 + idx * 1.68)
        box = s6.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, Inches(2.2), Inches(1.5), Inches(3.2))
        box.fill.solid()
        box.fill.fore_color.rgb = COLOR_CARD
        box.line.color.rgb = COLOR_BLUE_SUB
        box.line.width = Pt(1.5)

        tf = box.text_frame
        tf.word_wrap = True
        tf.margin_left = Inches(0.1)
        tf.margin_right = Inches(0.1)

        p = tf.paragraphs[0]
        p.text = step_title
        p.font.size = Pt(11)
        p.font.bold = True
        p.font.color.rgb = COLOR_BLUE_HEADING

        p2 = tf.add_paragraph()
        p2.text = step_desc
        p2.font.size = Pt(9.5)
        p2.font.color.rgb = COLOR_DARK_TEXT
        p2.space_before = Pt(8)

        if idx < len(steps_s6) - 1:
            arr = s6.shapes.add_shape(MSO_SHAPE.RIGHT_ARROW, Inches(0.8 + (idx+1)*1.68 - 0.22), Inches(3.5), Inches(0.18), Inches(0.3))
            arr.fill.solid()
            arr.fill.fore_color.rgb = COLOR_BLUE_SUB
            arr.line.fill.background()

    add_notes(s6, "This flowchart details the 7 sequential stages from raw document ingestion to vector matching and recommendation generation.")

    # ==========================================
    # SLIDE 7: SYSTEM ARCHITECTURE FLOWCHART
    # ==========================================
    s7 = prs.slides.add_slide(blank_slide_layout)
    set_slide_background(s7)
    add_header(s7, "Layered System Architecture Flowchart")

    layers = [
        ("USER INTERFACE LAYER", "React.js Frontend | Tailwind CSS | Recharts Dashboard", COLOR_BLUE_SUB),
        ("API & CONTROLLER LAYER", "FastAPI REST Server | CORS Middleware | PII Masking Engine", COLOR_TEAL),
        ("NLP & ML CORE ENGINE", "PyMuPDF Text Extractor | TF-IDF Vectorizer | Cosine Matcher | Skill Normalizer", COLOR_BLUE_HEADING),
        ("PERSISTENCE & DATA LAYER", "SQLite Database | SQLAlchemy 2.0 ORM | Local Skill Dictionary", COLOR_SUBTEXT)
    ]

    for idx, (layer_title, layer_desc, color) in enumerate(layers):
        top = Inches(1.5 + idx * 1.35)
        box = s7.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(1.5), top, Inches(10.333), Inches(1.1))
        box.fill.solid()
        box.fill.fore_color.rgb = COLOR_CARD
        box.line.color.rgb = color
        box.line.width = Pt(2)

        tf = box.text_frame
        tf.word_wrap = True
        tf.margin_left = Inches(0.3)
        tf.margin_top = Inches(0.15)

        p = tf.paragraphs[0]
        p.text = layer_title
        p.font.size = Pt(13)
        p.font.bold = True
        p.font.color.rgb = color

        p2 = tf.add_paragraph()
        p2.text = layer_desc
        p2.font.size = Pt(11)
        p2.font.color.rgb = COLOR_DARK_TEXT
        p2.space_before = Pt(4)

    add_notes(s7, "Layered architecture cleanly separates frontend presentation, FastAPI routing, core ML processing, and SQLite storage.")

    # ==========================================
    # SLIDE 8: TEXT PREPROCESSING FLOWCHART
    # ==========================================
    s8 = prs.slides.add_slide(blank_slide_layout)
    set_slide_background(s8)
    add_header(s8, "NLP Text Preprocessing Pipeline Flowchart")

    prep_steps = [
        ("Raw Text Extraction", "PyMuPDF/Docx extract text strings"),
        ("Text Cleaning", "Remove special symbols, URLs, email addresses"),
        ("Technical Token Protection", "Preserve C++, C#, .NET, React.js tokens"),
        ("Case Normalization", "Convert text to lowercase for uniformity"),
        ("Stop-word Filtering", "Remove common English words (and, the, with)"),
        ("TF-IDF Matrix", "Generate sparse numerical feature vectors")
    ]

    for idx, (title, desc) in enumerate(prep_steps):
        row = idx // 3
        col = idx % 3
        left = Inches(0.8 + col * 3.95)
        top = Inches(1.6 + row * 2.6)

        card = s8.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, Inches(3.8), Inches(2.2))
        card.fill.solid()
        card.fill.fore_color.rgb = COLOR_CARD
        card.line.color.rgb = COLOR_CARD_BORDER

        tf = card.text_frame
        tf.word_wrap = True
        tf.margin_left = Inches(0.2)

        p = tf.paragraphs[0]
        p.text = f"Step {idx+1}: {title}"
        p.font.size = Pt(12)
        p.font.bold = True
        p.font.color.rgb = COLOR_BLUE_HEADING

        p2 = tf.add_paragraph()
        p2.text = desc
        p2.font.size = Pt(10.5)
        p2.font.color.rgb = COLOR_DARK_TEXT
        p2.space_before = Pt(6)

    add_notes(s8, "Preprocessing transforms noisy document text into clean, structured tokens ready for TF-IDF vectorization.")

    # ==========================================
    # SLIDE 9: TF-IDF & COSINE MATHEMATICS
    # ==========================================
    s9 = prs.slides.add_slide(blank_slide_layout)
    set_slide_background(s9)
    add_header(s9, "TF-IDF Vectorization & Cosine Similarity Mathematics")

    box_l = s9.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.5), Inches(5.6), Inches(5.3))
    box_l.fill.solid()
    box_l.fill.fore_color.rgb = COLOR_CARD
    box_l.line.color.rgb = COLOR_CARD_BORDER
    tf_l = box_l.text_frame
    tf_l.word_wrap = True
    tf_l.margin_left = Inches(0.3)
    tf_l.margin_top = Inches(0.3)

    p = tf_l.paragraphs[0]
    p.text = "TF-IDF Mathematical Formulation"
    p.font.size = Pt(15)
    p.font.bold = True
    p.font.color.rgb = COLOR_BLUE_HEADING

    tf_bullets = [
        "Term Frequency (TF): Measures how frequently a skill term t appears in document d.",
        "Inverse Document Frequency (IDF): Evaluates term rarity across overall corpus: IDF(t) = log(N / df(t)).",
        "TF-IDF Weight: W(t,d) = TF(t,d) * IDF(t). Gives high weights to specific technical skills like PyMuPDF."
    ]
    for b in tf_bullets:
        p = tf_l.add_paragraph()
        p.text = "• " + b
        p.font.size = Pt(11)
        p.font.color.rgb = COLOR_DARK_TEXT
        p.space_after = Pt(10)

    box_r = s9.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(6.8), Inches(1.5), Inches(5.733), Inches(5.3))
    box_r.fill.solid()
    box_r.fill.fore_color.rgb = COLOR_CARD
    box_r.line.color.rgb = COLOR_BLUE_SUB
    box_r.line.width = Pt(1.5)
    tf_r = box_r.text_frame
    tf_r.word_wrap = True
    tf_r.margin_left = Inches(0.3)
    tf_r.margin_top = Inches(0.3)

    p = tf_r.paragraphs[0]
    p.text = "Cosine Similarity Vector Alignment"
    p.font.size = Pt(15)
    p.font.bold = True
    p.font.color.rgb = COLOR_BLUE_HEADING

    p = tf_r.add_paragraph()
    p.text = "Cosine Similarity Formula:"
    p.font.size = Pt(12)
    p.font.bold = True
    p.font.color.rgb = COLOR_TEAL

    p = tf_r.add_paragraph()
    p.text = "Similarity(A, B) = (A · B) / (||A|| * ||B||)"
    p.font.size = Pt(14)
    p.font.bold = True
    p.font.color.rgb = COLOR_BLUE_HEADING
    p.space_after = Pt(12)

    cos_bullets = [
        "A = Resume TF-IDF Vector | B = Job Description TF-IDF Vector",
        "Computes the cosine of the angle between two n-dimensional vectors.",
        "Range: 0.0 (No overlap) to 1.0 (100% exact match). Length-invariant comparison."
    ]
    for b in cos_bullets:
        p = tf_r.add_paragraph()
        p.text = "• " + b
        p.font.size = Pt(11)
        p.font.color.rgb = COLOR_DARK_TEXT
        p.space_after = Pt(8)

    add_notes(s9, "TF-IDF converts text into sparse numerical vectors, and Cosine Similarity measures the angular distance between candidate skills and job requirements.")

    # ==========================================
    # SLIDE 10: SKILL EXTRACTION FLOWCHART
    # ==========================================
    s10 = prs.slides.add_slide(blank_slide_layout)
    set_slide_background(s10)
    add_header(s10, "Skill Extraction & Synonym Normalization Engine")

    c_ext = s10.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.5), Inches(11.733), Inches(5.3))
    c_ext.fill.solid()
    c_ext.fill.fore_color.rgb = COLOR_CARD
    c_ext.line.color.rgb = COLOR_CARD_BORDER
    tf_e = c_ext.text_frame
    tf_e.word_wrap = True
    tf_e.margin_left = Inches(0.4)
    tf_e.margin_top = Inches(0.3)

    p = tf_e.paragraphs[0]
    p.text = "Transparent Skill Dictionary & Alias Normalization"
    p.font.size = Pt(16)
    p.font.bold = True
    p.font.color.rgb = COLOR_BLUE_HEADING

    e_bullets = [
        "Curated Skill Catalog: 150+ standard CSE skills mapped across Languages, Frameworks, ML/AI, Databases, and Tools.",
        "Synonym & Alias Mapping: Normalizes variations (e.g., 'JS' -> 'JavaScript', 'React.js' -> 'React', 'Sklearn' -> 'Scikit-learn').",
        "Aho-Corasick & Regex Parsing: Fast keyword lookup preventing false positives during extraction.",
        "Categorization: Groups extracted skills into core domains for structured dashboard visualization."
    ]
    for b in e_bullets:
        p = tf_e.add_paragraph()
        p.text = "• " + b
        p.font.size = Pt(12)
        p.font.color.rgb = COLOR_DARK_TEXT
        p.space_after = Pt(12)

    add_notes(s10, "Skill normalization resolves alias variations so equivalent candidate qualifications are correctly matched.")

    # ==========================================
    # SLIDE 11: SKILL GAP & ROADMAP ENGINE
    # ==========================================
    s11 = prs.slides.add_slide(blank_slide_layout)
    set_slide_background(s11)
    add_header(s11, "Skill Gap Analysis & Recommendation Engine")

    c_g = s11.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.5), Inches(5.6), Inches(5.3))
    c_g.fill.solid()
    c_g.fill.fore_color.rgb = COLOR_CARD
    c_g.line.color.rgb = COLOR_CARD_BORDER
    tf_g = c_g.text_frame
    tf_g.word_wrap = True
    tf_g.margin_left = Inches(0.3)
    tf_g.margin_top = Inches(0.3)

    p = tf_g.paragraphs[0]
    p.text = "Skill Gap Identification"
    p.font.size = Pt(15)
    p.font.bold = True
    p.font.color.rgb = COLOR_BLUE_HEADING

    g_bullets = [
        "Set Difference Operation: Missing Skills = Job Required Skills - Resume Skills.",
        "Identify Key Weaknesses: Pinpoints exact technical gaps like Docker, PyTorch, or Spring Boot.",
        "Matched Skills Recognition: Highlights candidate strengths to boost interview confidence."
    ]
    for b in g_bullets:
        p = tf_g.add_paragraph()
        p.text = "• " + b
        p.font.size = Pt(11)
        p.font.color.rgb = COLOR_DARK_TEXT
        p.space_after = Pt(10)

    c_rec = s11.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(6.8), Inches(1.5), Inches(5.733), Inches(5.3))
    c_rec.fill.solid()
    c_rec.fill.fore_color.rgb = COLOR_CARD
    c_rec.line.color.rgb = COLOR_BLUE_SUB
    c_rec.line.width = Pt(1.5)
    tf_r = c_rec.text_frame
    tf_r.word_wrap = True
    tf_r.margin_left = Inches(0.3)
    tf_r.margin_top = Inches(0.3)

    p = tf_r.paragraphs[0]
    p.text = "Actionable Learning Recommendations"
    p.font.size = Pt(15)
    p.font.bold = True
    p.font.color.rgb = COLOR_BLUE_HEADING

    r_bullets = [
        "Interactive Roadmap: Provides actionable study modules for each identified missing skill.",
        "Status Tracking: Candidates toggle tasks between 'Not Started', 'In Progress', and 'Completed'.",
        "Career Assistance Tool: Focuses on candidate growth rather than automated rejection."
    ]
    for b in r_bullets:
        p = tf_r.add_paragraph()
        p.text = "• " + b
        p.font.size = Pt(11)
        p.font.color.rgb = COLOR_DARK_TEXT
        p.space_after = Pt(10)

    add_notes(s11, "Skill gap analysis turns match scores into constructive learning roadmaps for student skill enhancement.")

    # ==========================================
    # SLIDE 12: MODEL & ALGORITHM COMPARISON (TABLE 2 OF 4)
    # ==========================================
    s12 = prs.slides.add_slide(blank_slide_layout)
    set_slide_background(s12)
    add_header(s12, "Algorithm & Match Approach Comparison")

    t_shape12 = s12.shapes.add_table(5, 4, Inches(0.8), Inches(1.5), Inches(11.733), Inches(5.3))
    t12 = t_shape12.table
    t12.columns[0].width = Inches(2.5)
    t12.columns[1].width = Inches(3.0)
    t12.columns[2].width = Inches(3.0)
    t12.columns[3].width = Inches(3.233)

    headers12 = ["Approach / Model", "Matching Logic", "Advantages", "Limitations / Trade-offs"]
    for c_idx, h in enumerate(headers12):
        cell = t12.cell(0, c_idx)
        cell.fill.solid()
        cell.fill.fore_color.rgb = COLOR_CARD
        p = cell.text_frame.paragraphs[0]
        p.text = h
        p.font.size = Pt(11)
        p.font.bold = True
        p.font.color.rgb = COLOR_BLUE_HEADING

    rows12 = [
        ("Exact Keyword Search", "String matching", "Fast execution", "Ignores synonyms & context"),
        ("TF-IDF + Cosine (Chosen)", "Term frequency vector similarity", "Weighted term importance, fast & transparent", "Requires dictionary for synonym resolution"),
        ("Word2Vec / GloVe", "Dense vector embeddings", "Captures semantic context", "Higher computational overhead"),
        ("Transformer (SBERT)", "Deep contextual embeddings", "State-of-the-art semantic match", "Heavy memory requirement & latency")
    ]

    for r_idx, r_data in enumerate(rows12):
        for c_idx, val in enumerate(r_data):
            cell = t12.cell(r_idx+1, c_idx)
            cell.fill.solid()
            cell.fill.fore_color.rgb = COLOR_BG
            p = cell.text_frame.paragraphs[0]
            p.text = val
            p.font.size = Pt(10.5)
            p.font.color.rgb = COLOR_DARK_TEXT

    apply_thick_black_borders(t12)
    add_notes(s12, "TF-IDF with Cosine Similarity provides the optimal balance of execution speed, term weighting, and transparent scoring for real-time applications.")

    # ==========================================
    # SLIDE 13: PROJECT IMPLEMENTATION
    # ==========================================
    s13 = prs.slides.add_slide(blank_slide_layout)
    set_slide_background(s13)
    add_header(s13, "Project Implementation & Core Software Modules")

    mods = [
        ("FastAPI Backend", "Python 3.14 REST server handling file uploads, preprocessing, similarity scoring, and SQLite persistence."),
        ("React 18 Frontend", "Vite-powered SPA with Tailwind CSS responsive dashboard, interactive charts, and dual-pane text editor."),
        ("Document Parsers", "PyMuPDF for multi-page PDF text extraction; python-docx for DOCX document parsing."),
        ("SQLite Database", "SQLAlchemy 2.0 ORM storing historical analysis records, similarity scores, and timestamps.")
    ]

    for idx, (title, desc) in enumerate(mods):
        row = idx // 2
        col = idx % 2
        left = Inches(0.8 + col * 5.95)
        top = Inches(1.6 + row * 2.6)

        card = s13.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, Inches(5.75), Inches(2.2))
        card.fill.solid()
        card.fill.fore_color.rgb = COLOR_CARD
        card.line.color.rgb = COLOR_CARD_BORDER

        tf = card.text_frame
        tf.word_wrap = True
        tf.margin_left = Inches(0.3)

        p = tf.paragraphs[0]
        p.text = title
        p.font.size = Pt(14)
        p.font.bold = True
        p.font.color.rgb = COLOR_BLUE_HEADING

        p2 = tf.add_paragraph()
        p2.text = desc
        p2.font.size = Pt(11)
        p2.font.color.rgb = COLOR_DARK_TEXT
        p2.space_before = Pt(6)

    add_notes(s13, "Full-stack implementation combines Python FastAPI backend performance with React frontend interactivity.")

    # ==========================================
    # SLIDE 14: FEATURE MATRIX (TABLE 3 OF 4)
    # ==========================================
    s14 = prs.slides.add_slide(blank_slide_layout)
    set_slide_background(s14)
    add_header(s14, "System Architecture Feature Matrix")

    t_shape14 = s14.shapes.add_table(6, 3, Inches(0.8), Inches(1.5), Inches(11.733), Inches(5.3))
    t14 = t_shape14.table
    t14.columns[0].width = Inches(3.5)
    t14.columns[1].width = Inches(4.5)
    t14.columns[2].width = Inches(3.733)

    headers14 = ["System Module", "Technical Technology / Tool", "Implementation Function"]
    for c_idx, h in enumerate(headers14):
        cell = t14.cell(0, c_idx)
        cell.fill.solid()
        cell.fill.fore_color.rgb = COLOR_CARD
        p = cell.text_frame.paragraphs[0]
        p.text = h
        p.font.size = Pt(11)
        p.font.bold = True
        p.font.color.rgb = COLOR_BLUE_HEADING

    rows14 = [
        ("Document Ingestion", "PyMuPDF, python-docx", "Multi-format text extraction from PDF/DOCX files"),
        ("Text Preprocessing", "Regex, NLTK stop-words", "Text cleaning, tokenization, synonym normalization"),
        ("Feature Vectorization", "Scikit-Learn TfidfVectorizer", "Numerical sparse feature matrix generation"),
        ("Similarity Matching", "Scikit-Learn cosine_similarity", "Vector dot-product calculation for score"),
        ("Data Persistence", "SQLite, SQLAlchemy 2.0", "Historical analysis record storage")
    ]

    for r_idx, r_data in enumerate(rows14):
        for c_idx, val in enumerate(r_data):
            cell = t14.cell(r_idx+1, c_idx)
            cell.fill.solid()
            cell.fill.fore_color.rgb = COLOR_BG
            p = cell.text_frame.paragraphs[0]
            p.text = val
            p.font.size = Pt(10.5)
            p.font.color.rgb = COLOR_DARK_TEXT

    apply_thick_black_borders(t14)
    add_notes(s14, "Feature matrix summarizes core system modules and corresponding technology stack choices.")

    # ==========================================
    # SLIDE 15: CHALLENGES & SOLUTIONS
    # ==========================================
    s15 = prs.slides.add_slide(blank_slide_layout)
    set_slide_background(s15)
    add_header(s15, "Engineering Challenges & Practical Solutions")

    challs = [
        ("1. Unstructured Resume Formatting", "Challenge: PDF layouts caused jumbled text strings during extraction.\nSolution: Implemented PyMuPDF block position sorting and clean newline joining."),
        ("2. Punctuation Stripping Bug", "Challenge: Standard regex destroyed technical tokens like 'C++', 'C#', '.NET', 'React.js'.\nSolution: Added pre-tokenization regex mappings (C++ -> cplusplus, .NET -> dotnet)."),
        ("3. Synonym Variations", "Challenge: Candidates write 'JS' while job specs request 'JavaScript'.\nSolution: Created transparent JSON skill catalog mapping aliases to canonical skill names.")
    ]

    for idx, (title, desc) in enumerate(challs):
        top = Inches(1.5 + idx * 1.75)
        card = s15.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), top, Inches(11.733), Inches(1.5))
        card.fill.solid()
        card.fill.fore_color.rgb = COLOR_CARD
        card.line.color.rgb = COLOR_CARD_BORDER

        tf = card.text_frame
        tf.word_wrap = True
        tf.margin_left = Inches(0.3)
        tf.margin_top = Inches(0.15)

        p = tf.paragraphs[0]
        p.text = title
        p.font.size = Pt(13)
        p.font.bold = True
        p.font.color.rgb = COLOR_BLUE_HEADING

        p2 = tf.add_paragraph()
        p2.text = desc
        p2.font.size = Pt(10.5)
        p2.font.color.rgb = COLOR_DARK_TEXT
        p2.space_before = Pt(4)

    add_notes(s15, "These technical challenges arose during development and were resolved through custom preprocessing logic.")

    # ==========================================
    # SLIDE 16: MODEL EVALUATION (TABLE 4 OF 4)
    # ==========================================
    s16 = prs.slides.add_slide(blank_slide_layout)
    set_slide_background(s16)
    add_header(s16, "Model Evaluation & Verification Metrics")

    t_shape16 = s16.shapes.add_table(5, 4, Inches(0.8), Inches(1.5), Inches(11.733), Inches(5.3))
    t16 = t_shape16.table
    t16.columns[0].width = Inches(3.0)
    t16.columns[1].width = Inches(2.8)
    t16.columns[2].width = Inches(2.8)
    t16.columns[3].width = Inches(3.133)

    headers16 = ["Test Scenario", "Expected Outcome", "Observed Result", "System Status"]
    for c_idx, h in enumerate(headers16):
        cell = t16.cell(0, c_idx)
        cell.fill.solid()
        cell.fill.fore_color.rgb = COLOR_CARD
        p = cell.text_frame.paragraphs[0]
        p.text = h
        p.font.size = Pt(11)
        p.font.bold = True
        p.font.color.rgb = COLOR_BLUE_HEADING

    rows16 = [
        ("Identical Text Match", "Similarity = 100%", "Similarity = 100.0%", "PASSED"),
        ("Partial Skill Match", "Similarity ~ 60-80%", "Similarity = 74.2%", "PASSED"),
        ("Unrelated Domain Text", "Similarity < 20%", "Similarity = 8.5%", "PASSED"),
        ("Multi-page PDF Parsing", "Clean text extraction", "100% text recovered", "PASSED")
    ]

    for r_idx, r_data in enumerate(rows16):
        for c_idx, val in enumerate(r_data):
            cell = t16.cell(r_idx+1, c_idx)
            cell.fill.solid()
            cell.fill.fore_color.rgb = COLOR_BG
            p = cell.text_frame.paragraphs[0]
            p.text = val
            p.font.size = Pt(10.5)
            p.font.color.rgb = COLOR_DARK_TEXT if c_idx < 3 else COLOR_TEAL

    apply_thick_black_borders(t16)
    add_notes(s16, "System verification confirmed robust matching behavior across exact, partial, and unrelated document pairs.")

    # ==========================================
    # SLIDE 17: LIVE DEMO 1 (RESUME UPLOAD SCREENSHOT)
    # ==========================================
    s17 = prs.slides.add_slide(blank_slide_layout)
    set_slide_background(s17)
    add_header(s17, "Live Demo 1 — Resume Upload & Text Extraction Interface", "APPLICATION DEMO SCREENSHOTS")

    path_demo1 = r"C:\Users\geeth\.gemini\antigravity\scratch\demo_screenshots\demo_1_upload.png"
    if os.path.exists(path_demo1):
        s17.shapes.add_picture(path_demo1, Inches(0.8), Inches(1.5), Inches(11.733), Inches(5.3))

    add_notes(s17, "Live Demo Screenshot 1 shows the React upload screen supporting drag-and-drop PDF/DOCX resumes and instant text preview.")

    # ==========================================
    # SLIDE 18: LIVE DEMO 2 (MATCH DASHBOARD SCREENSHOT)
    # ==========================================
    s18 = prs.slides.add_slide(blank_slide_layout)
    set_slide_background(s18)
    add_header(s18, "Live Demo 2 — Dual-Pane Job Matching & TF-IDF Score UI", "APPLICATION DEMO SCREENSHOTS")

    path_demo2 = r"C:\Users\geeth\.gemini\antigravity\scratch\demo_screenshots\demo_2_analyze.png"
    if os.path.exists(path_demo2):
        s18.shapes.add_picture(path_demo2, Inches(0.8), Inches(1.5), Inches(11.733), Inches(5.3))

    add_notes(s18, "Live Demo Screenshot 2 displays the dual-pane editor showing match score percentage, matched skills badge, and missing skills breakdown.")

    # ==========================================
    # SLIDE 19: LIVE DEMO 3 (LEARNING ROADMAP SCREENSHOT)
    # ==========================================
    s19 = prs.slides.add_slide(blank_slide_layout)
    set_slide_background(s19)
    add_header(s19, "Live Demo 3 — Skill Gap & Interactive Learning Roadmap UI", "APPLICATION DEMO SCREENSHOTS")

    path_demo3 = r"C:\Users\geeth\.gemini\antigravity\scratch\demo_screenshots\demo_3_roadmap.png"
    if os.path.exists(path_demo3):
        s19.shapes.add_picture(path_demo3, Inches(0.8), Inches(1.5), Inches(11.733), Inches(5.3))

    add_notes(s19, "Live Demo Screenshot 3 demonstrates the interactive skill roadmap where students can toggle learning tasks between Not Started, In Progress, and Completed.")

    # ==========================================
    # SLIDE 20: LIVE DEMO 4 (HISTORY LOG SCREENSHOT)
    # ==========================================
    s20 = prs.slides.add_slide(blank_slide_layout)
    set_slide_background(s20)
    add_header(s20, "Live Demo 4 — SQLite Analysis History Log & Skill Catalog UI", "APPLICATION DEMO SCREENSHOTS")

    path_demo4 = r"C:\Users\geeth\.gemini\antigravity\scratch\demo_screenshots\demo_4_history.png"
    if os.path.exists(path_demo4):
        s20.shapes.add_picture(path_demo4, Inches(0.8), Inches(1.5), Inches(11.733), Inches(5.3))

    add_notes(s20, "Live Demo Screenshot 4 displays the SQLite history table storing prior analyses with timestamps and similarity scores. In conclusion, this internship project successfully bridges candidate resumes and job descriptions using NLP and ML. Thank you!")

    # Save Presentation
    downloads_path = r"C:\Users\geeth\Downloads\AI_Based_Resume_Screening_Internship_Presentation_Meda_Geethika_HUMANIZED_LIVE_DEMO.pptx"
    try:
        prs.save(downloads_path)
        print(f"Presentation saved to Downloads: {os.path.abspath(downloads_path)}")
    except PermissionError:
        downloads_path_v2 = r"C:\Users\geeth\Downloads\AI_Based_Resume_Screening_Internship_Presentation_Meda_Geethika_HUMANIZED_LIVE_DEMO_v2.pptx"
        prs.save(downloads_path_v2)
        print(f"Saved to Downloads (v2): {os.path.abspath(downloads_path_v2)}")

    workspace_path = os.path.join(os.path.dirname(__file__), "..", "AI_Based_Resume_Screening_Internship_Presentation_Meda_Geethika_HUMANIZED_LIVE_DEMO.pptx")
    try:
        prs.save(workspace_path)
        print(f"Presentation saved to workspace: {os.path.abspath(workspace_path)}")
    except PermissionError:
        pass

if __name__ == "__main__":
    create_humanized_ppt()
